import os
import re
import requests
import xlrd  
from flask import Flask, request, jsonify
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

app = Flask(__name__)

BACKEND_URL = os.getenv("BACKEND_URL", "http://localhost:8080")
UPLOAD_FOLDER = "/app/uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

DNI_PATTERN = re.compile(r"DNI\s*(\d+)", re.IGNORECASE)
RD_PATTERN = re.compile(r"^(RD|R\.D\.|R\.M\.|RM|DS|LEY|DL)\s*[\d\-\./]+", re.IGNORECASE)
IGNORAR_CONCEPTOS = {"REINTEGRO", "ESCOLARIDAD", "AGUINALDO", "DETALLE", "TOTAL", "SUBTOTAL"}


def parsear_valor(v):
    if v is None:
        return None
    try:
        f = float(v)
        return round(f, 2) if f else None
    except (TypeError, ValueError):
        return None


def limpiar_concepto(c: str) -> str:
    return c.strip().lstrip("+").strip().upper()


def extraer_dni(texto: str):
    if not texto:
        return None
    m = DNI_PATTERN.search(str(texto))
    return m.group(1) if m else None


def _leer_filas(filepath: str) -> list:
    """Return all rows as a list of tuples, supporting both .xls and .xlsx."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".xls":
        wb = xlrd.open_workbook(filepath)
        ws = wb.sheet_by_index(0)
        filas = []
        for rx in range(ws.nrows):
            row = []
            for cx in range(ws.ncols):
                cell = ws.cell(rx, cx)
                # xlrd type 0=empty,1=text,2=number,3=date,4=bool,5=error,6=blank
                if cell.ctype in (0, 5, 6):
                    row.append(None)
                elif cell.ctype == 2:
                    # Return int if value is whole number, else float
                    v = cell.value
                    row.append(int(v) if v == int(v) else v)
                else:
                    row.append(cell.value)
            filas.append(tuple(row))
        return filas
    else:
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        return list(ws.iter_rows(values_only=True))


def extraer_empleados(filepath: str) -> list:
    """
    Parse a payroll Excel with vertical-block layout:

      Col A    | Col B           | Col C    | Col D
      ---------+-----------------+----------+--------
      HABERES  | Apellidos Nomb. | BASICA   | 0.03
      (merged) | (merged)        | DL19990  | 60.00
                | CARGO/PUESTO    | TPH      | 19.20
                | RD 150-93       | ...
                | uu-01-0-005     |
      DSCTOS   |                 | DL20530  | 3.80
      TOTAL HABERES              |          | 153.92
      TOTAL DESCUENTOS           |          | 76.27
      TOTAL LIQUIDO              |          | 77.65

    Mes/Año are NOT read from the file; they come from the user request.
    """
    filas = _leer_filas(filepath)

    empleados = []
    i = 0
    n = len(filas)

    while i < n:
        fila = filas[i]
        col_a = str(fila[0]).strip() if fila[0] else ""

        if col_a == "HABERES" and len(fila) > 1 and fila[1]:
            emp = {
                "nombre": str(fila[1]).strip(),
                "cargo": None,
                "resolucion": None,
                "codigo": None,
                "dni": None,
                "haberes": [],
                "descuentos": [],
                "total_haberes": None,
                "total_descuentos": None,
                "total_liquido": None,
            }

            concepto_inicial = str(fila[2]).strip() if len(fila) > 2 and fila[2] else ""
            valor_inicial = parsear_valor(fila[3] if len(fila) > 3 else None)
            if concepto_inicial and limpiar_concepto(concepto_inicial) not in IGNORAR_CONCEPTOS:
                if valor_inicial is not None:
                    emp["haberes"].append({"concepto": concepto_inicial, "monto": valor_inicial})

            seccion = "HABERES"
            i += 1

            while i < n:
                f = filas[i]
                a = str(f[0]).strip() if len(f) > 0 and f[0] else ""
                b = str(f[1]).strip() if len(f) > 1 and f[1] else ""
                c_cell = str(f[2]).strip() if len(f) > 2 and f[2] else ""
                d = parsear_valor(f[3] if len(f) > 3 else None)

                if a == "TOTAL HABERES":
                    emp["total_haberes"] = d
                    i += 1
                    continue
                if a == "TOTAL DESCUENTOS":
                    emp["total_descuentos"] = d
                    i += 1
                    continue
                if a == "TOTAL LIQUIDO":
                    emp["total_liquido"] = d
                    i += 1
                    break

                if a == "DSCTOS":
                    seccion = "DSCTOS"
                    if c_cell and limpiar_concepto(c_cell) not in IGNORAR_CONCEPTOS and d is not None:
                        emp["descuentos"].append({"concepto": c_cell, "monto": d})
                    i += 1
                    continue

                if a == "HABERES" and len(f) > 1 and f[1]:
                    break

                if b:
                    dni_encontrado = extraer_dni(b)
                    if dni_encontrado:
                        emp["dni"] = dni_encontrado
                    elif RD_PATTERN.match(b):
                        if not emp["resolucion"]:
                            emp["resolucion"] = b.strip()
                    elif re.match(r"^uu-", b, re.IGNORECASE):
                        emp["codigo"] = b
                    elif not re.match(r"^(TOTAL|DSCTOS|CARGO|PUESTO|DNI|RD)", b, re.IGNORECASE):
                        if emp["cargo"] is None:
                            emp["cargo"] = b

                if c_cell:
                    nombre_limpio = limpiar_concepto(c_cell)
                    if nombre_limpio not in IGNORAR_CONCEPTOS and not nombre_limpio.startswith("+"):
                        if d is not None:
                            item = {"concepto": c_cell.strip(), "monto": d}
                            if seccion == "HABERES":
                                emp["haberes"].append(item)
                            else:
                                emp["descuentos"].append(item)

                i += 1

            empleados.append(emp)
            continue

        i += 1

    return empleados


def detectar_duplicados(empleados: list) -> dict:
    """Detecta duplicados por nombre o DNI en la lista de empleados."""
    por_dni = {}
    por_nombre = {}
    duplicados = []

    for emp in empleados:
        nombre_normalizado = emp.get("nombre", "").strip().upper()
        dni = emp.get("dni")

        if dni:
            if dni in por_dni:
                duplicados.append({
                    "tipo": "dni",
                    "dni": dni,
                    "nombres": [por_dni[dni], emp.get("nombre")]
                })
            else:
                por_dni[dni] = emp.get("nombre")

        if nombre_normalizado:
            if nombre_normalizado in por_nombre:
                duplicados.append({
                    "tipo": "nombre",
                    "nombre": emp.get("nombre"),
                    "dni": dni
                })
            else:
                por_nombre[nombre_normalizado] = emp.get("nombre")

    return {
        "duplicados": duplicados,
        "total_duplicados": len(duplicados),
        "por_dni": list(por_dni.keys()),
        "por_nombre": list(por_nombre.keys())
    }


def detectar_multiples_por_trabajador(empleados: list) -> dict:
    """Detecta si hay múltiples planillas para el mismo trabajador con不同的 neto."""
    por_dni = {}
    por_nombre_key = {}

    for emp in empleados:
        dni = emp.get("dni")
        nombre = emp.get("nombre", "").strip().upper()
        liquido = emp.get("total_liquido")

        key_dni = dni if dni else f"NOMBRE_{nombre}"
        key_nombre = f"{nombre}_{dni or 'sin_dni'}"

        if key_dni not in por_dni:
            por_dni[key_dni] = []

        por_dni[key_dni].append({
            "nombre_original": emp.get("nombre"),
            "nombre": nombre,
            "dni": dni,
            "liquido": liquido,
            "nombre_key": key_nombre,
            "haberes": emp.get("haberes", []),
            "descuentos": emp.get("descuentos", [])
        })

    multiples = []
    checked = set()

    for key, items in por_dni.items():
        if len(items) > 1:
            first = items[0]
            first_liq = first["liquido"] if first["liquido"] else 0
            all_same = all(abs((i["liquido"] or 0) - first_liq) < 0.01 for i in items)

            if not all_same:
                for item in items:
                    item_key = f"{item['nombre']}_{item['dni']}_{item['liquido']}"
                    if item_key not in checked:
                        checked.add(item_key)
                        multiples.append({
                            "nombre": item["nombre_original"],
                            "dni": item["dni"],
                            "liquido": item["liquido"],
                            "count": len(items),
                            "diferencia": abs(first_liq - (item["liquido"] or 0))
                        })

    return {
        "multiples": multiples,
        "total_multiples": len(multiples),
        "has_multiples": len(multiples) > 0
    }


# ─────────────────────────────────────────────────────────────────────────────

@app.route("/health", methods=["GET"])
def health():
    return jsonify({"status": "ok"})


@app.route("/process-excel", methods=["POST"])
def process_excel():
    if "file" not in request.files:
        return jsonify({"error": "No se encontró archivo"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Archivo sin nombre"}), 400

    filename = secure_filename(file.filename)
    if not filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Solo se aceptan archivos Excel (.xlsx, .xls)"}), 400

    mes = request.form.get("mes", type=int)
    anio = request.form.get("anio", type=int)

    if not mes or not anio:
        return jsonify({"error": "Se requieren los campos mes y anio"}), 400
    if not (1 <= mes <= 12):
        return jsonify({"error": "Mes debe estar entre 1 y 12"}), 400

    filepath = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    file.save(filepath)

    try:
        empleados = extraer_empleados(filepath)

        info_duplicados = detectar_duplicados(empleados)

        payload = {
            "mes": mes,
            "anio": anio,
            "total_empleados": len(empleados),
            "empleados": empleados,
            "verificar_duplicados": True,
        }

        response = requests.post(
            f"{BACKEND_URL}/api/importar/haberes",
            json=payload,
            timeout=60,
        )

        if response.status_code == 200:
            resp_data = response.json()
            return jsonify({
                "message": "Excel procesado correctamente",
                "personal_creados": resp_data.get("personal_creados", 0),
                "planillas_creadas": resp_data.get("planillas_creadas", 0),
                "personal": len(empleados),
                "planillas": len(empleados),
                "errores": resp_data.get("errores", []),
                "duplicados": info_duplicados.get("duplicados", []),
                "total_duplicados": info_duplicados.get("total_duplicados", 0),
            })
        else:
            return jsonify({
                "error": "Error al enviar al backend",
                "details": response.text,
            }), response.status_code

    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if os.path.exists(filepath):
            os.remove(filepath)


@app.route("/validate-excel", methods=["POST"])
def validate_excel():
    if "file" not in request.files:
        return jsonify({"error": "No se encontró archivo"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Archivo sin nombre"}), 400

    filename = secure_filename(file.filename)
    if not filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Solo se aceptan archivos Excel (.xlsx, .xls)"}), 400

    mes = request.form.get("mes", type=int)
    anio = request.form.get("anio", type=int)

    filepath = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    file.save(filepath)

    try:
        empleados = extraer_empleados(filepath)
        info_duplicados = detectar_duplicados(empleados)
        info_multiples = detectar_multiples_por_trabajador(empleados)
        return jsonify({
            "valid": True,
            "total_empleados": len(empleados),
            "preview": empleados[:5],
            "empleados": empleados,
            "duplicados": info_duplicados.get("duplicados", []),
            "total_duplicados": info_duplicados.get("total_duplicados", 0),
            "multiples": info_multiples.get("multiples", []),
            "total_multiples": info_multiples.get("total_multiples", 0),
            "has_multiples": info_multiples.get("has_multiples", False),
            "mes": mes,
            "anio": anio,
        })
    except Exception as e:
        return jsonify({"valid": False, "error": str(e)}), 400
    finally:
        if os.path.exists(filepath):
            os.remove(filepath)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8081)
